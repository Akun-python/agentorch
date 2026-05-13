from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BASE_DIR = Path(__file__).resolve().parent
SUMMARY_CSV = BASE_DIR / "baseline_reference_summary.csv"
OUTPUT_XLSX = BASE_DIR / "长期记忆基线共同指标整理.xlsx"


# 这里的状态含义是“能否直接作为当前统一对比表的同口径证据”。
STATUS_DIRECT = "直接报告"
STATUS_PARTIAL = "部分报告"
STATUS_NONE = "未直接报告"


BASELINE_ROWS = [
    {
        "baseline": "rag_chunk_memory",
        "method_source_ids": ["002"],
        "metric_source_ids": [],
        "single_hop": STATUS_NONE,
        "multi_hop": STATUS_NONE,
        "temporal": STATUS_NONE,
        "open_domain": STATUS_NONE,
        "overall": STATUS_NONE,
        "llm_judge": STATUS_NONE,
        "f1": STATUS_NONE,
        "bleu": STATUS_NONE,
        "token_context": STATUS_NONE,
        "suitable_main_table": "否",
        "note": "RAG 原论文更适合作为方法出处，不是当前 LoCoMo 五类统一主表的直接数值来源。",
    },
    {
        "baseline": "graph_rag_memory",
        "method_source_ids": ["003"],
        "metric_source_ids": [],
        "single_hop": STATUS_NONE,
        "multi_hop": STATUS_NONE,
        "temporal": STATUS_NONE,
        "open_domain": STATUS_NONE,
        "overall": STATUS_NONE,
        "llm_judge": STATUS_NONE,
        "f1": STATUS_NONE,
        "bleu": STATUS_NONE,
        "token_context": STATUS_NONE,
        "suitable_main_table": "否",
        "note": "GraphRAG 文献在当前参考包中主要承担方法出处作用，不是统一 LoCoMo 五类主表的直接指标来源。",
    },
    {
        "baseline": "light_rag_memory",
        "method_source_ids": ["004"],
        "metric_source_ids": [],
        "single_hop": STATUS_NONE,
        "multi_hop": STATUS_NONE,
        "temporal": STATUS_NONE,
        "open_domain": STATUS_NONE,
        "overall": STATUS_NONE,
        "llm_judge": STATUS_NONE,
        "f1": STATUS_NONE,
        "bleu": STATUS_NONE,
        "token_context": STATUS_NONE,
        "suitable_main_table": "否",
        "note": "LightRAG 论文在当前包内不是 LoCoMo 五类统一对比表的直接来源，适合保留为方法参考。",
    },
    {
        "baseline": "hippo_rag2_memory",
        "method_source_ids": ["005"],
        "metric_source_ids": [],
        "single_hop": STATUS_NONE,
        "multi_hop": STATUS_NONE,
        "temporal": STATUS_NONE,
        "open_domain": STATUS_NONE,
        "overall": STATUS_NONE,
        "llm_judge": STATUS_NONE,
        "f1": STATUS_NONE,
        "bleu": STATUS_NONE,
        "token_context": STATUS_NONE,
        "suitable_main_table": "否",
        "note": "HippoRAG2 更适合作为方法出处；若要并入统一主表，需要你自己按本地协议复现实验。",
    },
    {
        "baseline": "hypergraph_rag_memory",
        "method_source_ids": ["006"],
        "metric_source_ids": [],
        "single_hop": STATUS_NONE,
        "multi_hop": STATUS_NONE,
        "temporal": STATUS_NONE,
        "open_domain": STATUS_NONE,
        "overall": STATUS_NONE,
        "llm_judge": STATUS_NONE,
        "f1": STATUS_NONE,
        "bleu": STATUS_NONE,
        "token_context": STATUS_NONE,
        "suitable_main_table": "否",
        "note": "HyperGraphRAG 当前在本地参考包中主要是方法出处，未整理出可直接并入统一 LoCoMo 主表的同口径指标。",
    },
    {
        "baseline": "openai_memory",
        "method_source_ids": ["007"],
        "metric_source_ids": ["014"],
        "single_hop": STATUS_DIRECT,
        "multi_hop": STATUS_DIRECT,
        "temporal": STATUS_DIRECT,
        "open_domain": STATUS_DIRECT,
        "overall": STATUS_DIRECT,
        "llm_judge": STATUS_DIRECT,
        "f1": STATUS_DIRECT,
        "bleu": STATUS_DIRECT,
        "token_context": STATUS_NONE,
        "suitable_main_table": "是",
        "note": "方法来源是官方产品说明；统一 LoCoMo 五类分项与 F1/BLEU/LLM Judge 指标直接看 014。",
    },
    {
        "baseline": "langmem_memory",
        "method_source_ids": ["008"],
        "metric_source_ids": ["014"],
        "single_hop": STATUS_DIRECT,
        "multi_hop": STATUS_DIRECT,
        "temporal": STATUS_DIRECT,
        "open_domain": STATUS_DIRECT,
        "overall": STATUS_DIRECT,
        "llm_judge": STATUS_DIRECT,
        "f1": STATUS_DIRECT,
        "bleu": STATUS_DIRECT,
        "token_context": STATUS_NONE,
        "suitable_main_table": "是",
        "note": "LangMem README 负责方法出处；统一五类分项与 F1/BLEU/LLM Judge 直接来源是 014。",
    },
    {
        "baseline": "zep_memory",
        "method_source_ids": ["009"],
        "metric_source_ids": ["014"],
        "single_hop": STATUS_DIRECT,
        "multi_hop": STATUS_DIRECT,
        "temporal": STATUS_DIRECT,
        "open_domain": STATUS_DIRECT,
        "overall": STATUS_DIRECT,
        "llm_judge": STATUS_DIRECT,
        "f1": STATUS_DIRECT,
        "bleu": STATUS_DIRECT,
        "token_context": STATUS_NONE,
        "suitable_main_table": "是",
        "note": "Zep 论文负责方法出处；统一 LoCoMo 五类指标在当前本地参考包里最直接的来源是 014。",
    },
    {
        "baseline": "amem_memory",
        "method_source_ids": ["010"],
        "metric_source_ids": ["010"],
        "single_hop": STATUS_DIRECT,
        "multi_hop": STATUS_DIRECT,
        "temporal": STATUS_DIRECT,
        "open_domain": STATUS_DIRECT,
        "overall": STATUS_DIRECT,
        "llm_judge": STATUS_NONE,
        "f1": STATUS_DIRECT,
        "bleu": STATUS_DIRECT,
        "token_context": STATUS_DIRECT,
        "suitable_main_table": "是",
        "note": "010 直接给出 LoCoMo 五类分项、Overall、F1、BLEU 和 Token 长度相关信息；未整理到统一 LLM Judge 口径。",
    },
    {
        "baseline": "mem0_memory",
        "method_source_ids": ["011"],
        "metric_source_ids": ["011"],
        "single_hop": STATUS_DIRECT,
        "multi_hop": STATUS_DIRECT,
        "temporal": STATUS_DIRECT,
        "open_domain": STATUS_DIRECT,
        "overall": STATUS_DIRECT,
        "llm_judge": STATUS_DIRECT,
        "f1": STATUS_DIRECT,
        "bleu": STATUS_DIRECT,
        "token_context": STATUS_DIRECT,
        "suitable_main_table": "是",
        "note": "011 是当前本地参考包里最核心的统一口径来源之一，LoCoMo 五类、F1、BLEU、LLM Judge 与 Token 效率都比较完整。",
    },
    {
        "baseline": "mem0_graph_memory",
        "method_source_ids": ["011"],
        "metric_source_ids": ["011"],
        "single_hop": STATUS_DIRECT,
        "multi_hop": STATUS_DIRECT,
        "temporal": STATUS_DIRECT,
        "open_domain": STATUS_DIRECT,
        "overall": STATUS_DIRECT,
        "llm_judge": STATUS_DIRECT,
        "f1": STATUS_DIRECT,
        "bleu": STATUS_DIRECT,
        "token_context": STATUS_DIRECT,
        "suitable_main_table": "是",
        "note": "与 mem0_memory 共用 011；其中 Mem0-Graph 行直接提供统一口径分项与效率指标。",
    },
    {
        "baseline": "mirix_memory",
        "method_source_ids": ["012"],
        "metric_source_ids": ["012"],
        "single_hop": STATUS_DIRECT,
        "multi_hop": STATUS_DIRECT,
        "temporal": STATUS_DIRECT,
        "open_domain": STATUS_DIRECT,
        "overall": STATUS_DIRECT,
        "llm_judge": STATUS_DIRECT,
        "f1": STATUS_NONE,
        "bleu": STATUS_NONE,
        "token_context": STATUS_NONE,
        "suitable_main_table": "是",
        "note": "012 直接给出 LoCoMo 五类与 LLM Judge；效率更偏存储压缩，不是统一 Token/Context Length 主表口径。",
    },
    {
        "baseline": "memobase_memory",
        "method_source_ids": ["013", "014"],
        "metric_source_ids": ["014"],
        "single_hop": STATUS_DIRECT,
        "multi_hop": STATUS_DIRECT,
        "temporal": STATUS_DIRECT,
        "open_domain": STATUS_DIRECT,
        "overall": STATUS_DIRECT,
        "llm_judge": STATUS_DIRECT,
        "f1": STATUS_DIRECT,
        "bleu": STATUS_DIRECT,
        "token_context": STATUS_NONE,
        "suitable_main_table": "是",
        "note": "013 是产品介绍页；014 才是 LoCoMo benchmark 直接结果来源，且给了 v0.0.32 / v0.0.37 两个版本。",
    },
    {
        "baseline": "memu_memory",
        "method_source_ids": ["015"],
        "metric_source_ids": ["015"],
        "single_hop": STATUS_NONE,
        "multi_hop": STATUS_NONE,
        "temporal": STATUS_NONE,
        "open_domain": STATUS_NONE,
        "overall": f"{STATUS_PARTIAL}(仅 README 中的 LoCoMo 平均准确率)",
        "llm_judge": STATUS_NONE,
        "f1": STATUS_NONE,
        "bleu": STATUS_NONE,
        "token_context": STATUS_DIRECT,
        "suitable_main_table": "否",
        "note": "015 README 提到 LoCoMo 平均准确率和 Token 成本下降，但没有五类分项主表，适合做系统说明，不适合直接并统一分项表。",
    },
    {
        "baseline": "memos_memory",
        "method_source_ids": ["016"],
        "metric_source_ids": ["016"],
        "single_hop": STATUS_DIRECT,
        "multi_hop": STATUS_DIRECT,
        "temporal": STATUS_DIRECT,
        "open_domain": STATUS_DIRECT,
        "overall": STATUS_DIRECT,
        "llm_judge": STATUS_DIRECT,
        "f1": STATUS_DIRECT,
        "bleu": STATUS_DIRECT,
        "token_context": STATUS_DIRECT,
        "suitable_main_table": "是",
        "note": "016 直接给出 LoCoMo 五类、Overall、LLM Judge、Overall F1 和 Tokens/Context Length，是可直接并主表的强来源。",
    },
]


METRIC_GUIDE_ROWS = [
    {
        "metric": "Single-Hop",
        "description": "单跳问答，通常对应单段事实定位或单次会话内直接检索。",
        "recommended": "是",
        "source_ids": ["001", "014", "011", "012", "016", "010"],
        "note": "这是当前长期记忆对比里最常见的统一分项之一。",
    },
    {
        "metric": "Multi-Hop",
        "description": "多跳问答，需要跨片段、跨会话或跨实体进行组合推理。",
        "recommended": "是",
        "source_ids": ["001", "014", "011", "012", "016", "010"],
        "note": "适合体现图结构和记忆治理机制的优势。",
    },
    {
        "metric": "Temporal",
        "description": "时序推理，考察事件先后、时间变更和时间关系一致性。",
        "recommended": "是",
        "source_ids": ["001", "014", "011", "012", "016", "010"],
        "note": "对长期记忆系统很关键，尤其适合区分普通 RAG 与记忆系统。",
    },
    {
        "metric": "Open Domain",
        "description": "开放域问题，通常结合长期会话记忆与外部常识/世界知识。",
        "recommended": "是",
        "source_ids": ["001", "014", "011", "012", "016", "010"],
        "note": "该分项在不同论文里有时更依赖底模知识，解释结果时要单独注明。",
    },
    {
        "metric": "Overall",
        "description": "总体表现，通常是各类问题的汇总分数或综合准确率。",
        "recommended": "是",
        "source_ids": ["014", "011", "012", "016", "010", "015"],
        "note": "若仅有 Overall，没有各分项，建议标成部分可比。",
    },
    {
        "metric": "LLM-as-a-Judge",
        "description": "由评审 LLM 判断回答是否与标准答案语义一致的指标。",
        "recommended": "是",
        "source_ids": ["014", "011", "012", "016"],
        "note": "当前本地参考包里最适合作为统一主表的主指标。",
    },
    {
        "metric": "F1",
        "description": "词级别重合指标，兼顾 precision 和 recall。",
        "recommended": "建议补充",
        "source_ids": ["014", "011", "016", "010"],
        "note": "适合做补充指标，但不宜单独替代语义一致性判断。",
    },
    {
        "metric": "BLEU-1",
        "description": "生成文本的词面重合指标，通常比 F1 更表层。",
        "recommended": "建议补充",
        "source_ids": ["014", "011", "016", "010"],
        "note": "在长答案任务里容易低估语义正确但措辞不同的回答。",
    },
    {
        "metric": "Token / Context Length",
        "description": "上下文长度、输入 Token 或相对 Token 成本等效率指标。",
        "recommended": "建议补充",
        "source_ids": ["011", "016", "010", "015"],
        "note": "适合支撑“性能-代价”分析，但口径要统一后再横向比较。",
    },
]


THIN_SIDE = Side(style="thin", color="D9DEE7")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUB_HEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
GRAY_FILL = PatternFill("solid", fgColor="F3F4F6")
DIRECT_FILL = PatternFill("solid", fgColor="E2F0D9")
PARTIAL_FILL = PatternFill("solid", fgColor="FFF2CC")
NONE_FILL = PatternFill("solid", fgColor="FCE4D6")


def load_reference_map() -> dict[str, dict[str, str]]:
    with SUMMARY_CSV.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return {row["id"]: row for row in reader}


def join_source_values(reference_map: dict[str, dict[str, str]], source_ids: list[str], key: str) -> str:
    values: list[str] = []
    for source_id in source_ids:
        row = reference_map.get(source_id)
        if row is None:
            continue
        value = (row.get(key) or "").strip()
        if value:
            values.append(value)
    return "\n".join(values)


def format_source_ids(source_ids: list[str]) -> str:
    return ", ".join(source_ids) if source_ids else ""


def apply_base_style(ws) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)


def set_column_widths(ws, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def fill_status_cell(cell) -> None:
    value = str(cell.value or "")
    if value.startswith(STATUS_DIRECT):
        cell.fill = DIRECT_FILL
    elif value.startswith(STATUS_PARTIAL):
        cell.fill = PARTIAL_FILL
    elif value.startswith(STATUS_NONE):
        cell.fill = NONE_FILL


def build_baseline_sheet(wb: Workbook, reference_map: dict[str, dict[str, str]]) -> None:
    ws = wb.active
    ws.title = "基线共同指标"
    headers = [
        "基线",
        "方法来源ID",
        "方法来源标题",
        "指标直接来源ID",
        "指标直接来源标题",
        "Single-Hop",
        "Multi-Hop",
        "Temporal",
        "Open Domain",
        "Overall",
        "LLM-as-a-Judge",
        "F1",
        "BLEU-1",
        "Token/Context Length",
        "适合直接并入统一主表",
        "说明",
        "方法来源URL",
        "指标来源URL",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for row in BASELINE_ROWS:
        method_ids = row["method_source_ids"]
        metric_ids = row["metric_source_ids"]
        ws.append(
            [
                row["baseline"],
                format_source_ids(method_ids),
                join_source_values(reference_map, method_ids, "title"),
                format_source_ids(metric_ids),
                join_source_values(reference_map, metric_ids, "title"),
                row["single_hop"],
                row["multi_hop"],
                row["temporal"],
                row["open_domain"],
                row["overall"],
                row["llm_judge"],
                row["f1"],
                row["bleu"],
                row["token_context"],
                row["suitable_main_table"],
                row["note"],
                join_source_values(reference_map, method_ids, "landing_url"),
                join_source_values(reference_map, metric_ids, "landing_url"),
            ]
        )

    apply_base_style(ws)
    set_column_widths(
        ws,
        {
            "A": 24,
            "B": 12,
            "C": 42,
            "D": 14,
            "E": 42,
            "F": 16,
            "G": 16,
            "H": 16,
            "I": 16,
            "J": 16,
            "K": 18,
            "L": 12,
            "M": 12,
            "N": 22,
            "O": 16,
            "P": 48,
            "Q": 40,
            "R": 40,
        },
    )

    for row in ws.iter_rows(min_row=2, max_col=15):
        for cell in row[5:14]:
            fill_status_cell(cell)
        if row[14].value == "是":
            row[14].fill = DIRECT_FILL
        else:
            row[14].fill = NONE_FILL


def build_metric_guide_sheet(wb: Workbook, reference_map: dict[str, dict[str, str]]) -> None:
    ws = wb.create_sheet("指标说明")
    ws.append(["状态说明", "含义"])
    ws.append([STATUS_DIRECT, "当前来源中直接给出可并入统一对比表的同口径指标。"])
    ws.append([STATUS_PARTIAL, "当前来源只给出总体或相近指标，不能与五类分项完全对齐。"])
    ws.append([STATUS_NONE, "当前来源更适合作为方法出处，不是统一指标直引来源。"])
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for cell in ws["A"]:
        if cell.row >= 2:
            fill_status_cell(cell)

    ws.append([])
    ws.append(["指标", "中文说明", "建议作为主表", "主要直接来源ID", "主要直接来源标题", "备注"])
    for cell in ws[6]:
        cell.fill = SUB_HEADER_FILL
        cell.font = Font(bold=True, color="000000")

    for row in METRIC_GUIDE_ROWS:
        source_ids = row["source_ids"]
        ws.append(
            [
                row["metric"],
                row["description"],
                row["recommended"],
                format_source_ids(source_ids),
                join_source_values(reference_map, source_ids, "title"),
                row["note"],
            ]
        )

    apply_base_style(ws)
    set_column_widths(
        ws,
        {
            "A": 22,
            "B": 48,
            "C": 16,
            "D": 20,
            "E": 50,
            "F": 42,
        },
    )


def build_reference_sheet(wb: Workbook, reference_map: dict[str, dict[str, str]]) -> None:
    ws = wb.create_sheet("来源清单")
    headers = [
        "ID",
        "methods",
        "title",
        "kind",
        "status",
        "file",
        "landing_url",
        "source_url",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for source_id in sorted(reference_map):
        row = reference_map[source_id]
        ws.append([row.get(key, "") for key in ["id", "methods", "title", "kind", "status", "file", "landing_url", "source_url"]])

    apply_base_style(ws)
    set_column_widths(
        ws,
        {
            "A": 10,
            "B": 28,
            "C": 48,
            "D": 20,
            "E": 14,
            "F": 34,
            "G": 40,
            "H": 40,
        },
    )


def build_workbook() -> Path:
    reference_map = load_reference_map()
    wb = Workbook()
    wb.properties.creator = "Codex"
    wb.properties.title = "长期记忆基线共同指标整理"
    wb.properties.subject = "long_term_memory_graph baseline references"
    wb.properties.description = "基于本地参考文献包整理的对比基线共同指标说明。"

    build_baseline_sheet(wb, reference_map)
    build_metric_guide_sheet(wb, reference_map)
    build_reference_sheet(wb, reference_map)

    wb.save(OUTPUT_XLSX)
    return OUTPUT_XLSX


if __name__ == "__main__":
    output_path = build_workbook()
    print(output_path)
